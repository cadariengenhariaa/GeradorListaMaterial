"""Montagem das abas da lista de material (padrão das listas A056 e 648).

As abas são montadas num arquivo temporário do openpyxl; depois o módulo
pacote.py transfere essas abas para dentro do modelo da folha de rosto.
"""
import io
import math
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as letra
from openpyxl.worksheet.pagebreak import Break, RowBreak

from . import config
from .modelo import COMUM, NOME_ABA, PRIVATIVA, lancar_linha

FINO = Side(style="thin")
MEDIO = Side(style="medium")
SEM = Side(style=None)
EMU = 9525


def _fonte(tam=None, negrito=False, cor=None):
    return Font(name=config.FONTE, size=tam or config.TAM_FONTE, bold=negrito, color=cor)


def _preench(cor):
    return PatternFill("solid", start_color=cor, end_color=cor) if cor else PatternFill()


def _centro(quebra=True, h="center", v="center"):
    return Alignment(horizontal=h, vertical=v, wrap_text=quebra)


@dataclass
class AbaGerada:
    titulo: str
    aba: str
    ultima_coluna: int
    ultima_linha: int
    valores_formulas: dict = field(default_factory=dict)   # coordenada -> valor calculado
    itens: int = 0
    comentarios: int = 0
    mesclagens: int = 0


@dataclass
class DadosCapa:
    ca: str = ""
    obra: str = ""
    sistema: str = "CLIMATIZAÇÃO"
    etapa: str = "PROJETO EXECUTIVO"
    documento: str = "LISTA DE MATERIAIS"
    numero_documento: str = ""
    rev: str = "0"
    te: str = "A"
    descricao_rev: str = "EMISSÃO INICIAL"
    elab: str = ""
    verif: str = ""
    resp: str = ""
    data: object = None          # datetime.date
    desenhos: str = ""
    prefixo: str = "CLI"


def _largura_ambiente(titulo, unica_no_nivel=False, titulo_nivel=""):
    palavras = titulo.replace("\n", " ").split(" ")
    w = max(len(p) for p in palavras) if palavras else 4
    largura = max(15.0, min(30.0, w * 2.1 + 3))
    if len(titulo) > 16:
        largura = max(largura, 22.0)
    if unica_no_nivel:
        largura = max(largura, min(40.0, len(titulo_nivel) * 1.9 + 4))
    return largura


def _altura_descricao(txt):
    linhas = max(1, math.ceil(len(txt) / 82))
    return max(24.95, linhas * 20 + 6)


def _borda(col, ultima, top=FINO, bottom=FINO, left=None, right=None):
    esquerda = left if left is not None else (MEDIO if col == 1 else FINO)
    direita = right if right is not None else (MEDIO if col == ultima else FINO)
    return Border(left=esquerda, right=direita, top=top, bottom=bottom)


def montar_listas(analise, capa, logo_cadari=None, progresso=None):
    """Monta as abas e devolve (bytes do xlsx temporário, [AbaGerada])."""
    wb = Workbook()
    wb.remove(wb.active)
    colunas = analise.colunas()
    estrutura = analise.itens()
    geradas = []
    abas = [a for a in (COMUM, PRIVATIVA) if colunas[a]]
    for n_aba, aba in enumerate(abas):
        titulo = f"{capa.prefixo.strip() or 'CLI'} - {NOME_ABA[aba]}"[:31]
        ws = wb.create_sheet(titulo)
        g = _montar_aba(ws, analise, aba, colunas[aba], estrutura, capa, logo_cadari)
        if g.itens:
            geradas.append(g)
        else:
            wb.remove(ws)
        if progresso:
            progresso((n_aba + 1) / max(1, len(abas)))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), geradas


def _montar_aba(ws, analise, aba, colunas, estrutura, capa, logo_cadari):
    PRIMEIRA = 6
    n = len(colunas)
    ULT = PRIMEIRA + n - 1
    if ULT < PRIMEIRA + 2:
        ULT = PRIMEIRA + 2             # bloco do cabeçalho precisa de pelo menos 3 colunas
    L = letra(ULT)
    g = AbaGerada(titulo=ws.title, aba=aba, ultima_coluna=ULT, ultima_linha=0)
    valores = g.valores_formulas

    # ---------------- larguras ----------------
    for col, w in config.LARGURAS_FIXAS.items():
        ws.column_dimensions[col].width = w
    por_nivel = {}
    for c in colunas:
        por_nivel[c.nivel] = por_nivel.get(c.nivel, 0) + 1
    for i, c in enumerate(colunas):
        cab = analise.niveis[c.nivel].cabecalho
        ws.column_dimensions[letra(PRIMEIRA + i)].width = _largura_ambiente(c.titulo, por_nivel[c.nivel] == 1, cab)
    for i in range(n, ULT - PRIMEIRA + 1):
        ws.column_dimensions[letra(PRIMEIRA + i)].width = 15.0

    # ---------------- bloco do documento (linhas 1 a 9) ----------------
    for r in range(1, 10):
        for c in range(1, ULT + 1):
            cel = ws.cell(r, c)
            topo = MEDIO if r == 1 else (FINO if r in (5, 9) else SEM)
            base = FINO if r in (4, 8, 9) else SEM
            esq = MEDIO if c == 1 else SEM
            dir_ = MEDIO if c == ULT else SEM
            if r <= 4:
                if c in (2, 5):
                    dir_ = FINO
                if c == PRIMEIRA or c == ULT:
                    esq = FINO
                if c == ULT and r == 3:
                    topo = FINO
            cel.border = Border(left=esq, right=dir_, top=topo, bottom=base)
            cel.font = _fonte(negrito=True)
            cel.alignment = Alignment(vertical="center")
    ws.merge_cells("A1:B4")
    ws.merge_cells("C1:C4")
    ws.merge_cells(start_row=1, start_column=PRIMEIRA, end_row=1, end_column=ULT - 2)
    ws.merge_cells(start_row=2, start_column=PRIMEIRA, end_row=4, end_column=ULT - 2)
    capa_ref = f"'{config.ABA_CAPA}'"
    ws["A1"] = f"={capa_ref}!A1"
    ws["A1"].font = _fonte(negrito=True, cor="FF7F7F7F")
    ws["A1"].alignment = _centro(False)
    cel = ws.cell(1, PRIMEIRA, "Nº Documento")
    cel.font = _fonte(14)
    cel.alignment = Alignment(horizontal="center", vertical="top")
    cel = ws.cell(2, PRIMEIRA, '=MID(CELL("filename"),SEARCH("[",CELL("filename"))+1,'
                               'SEARCH("]",CELL("filename"))-SEARCH("[",CELL("filename"))-6)')
    cel.alignment = _centro(False)
    valores[cel.coordinate] = capa.numero_documento
    cel = ws.cell(1, ULT, f"={capa_ref}!L1")
    cel.font = _fonte()
    cel.alignment = Alignment(horizontal="center")
    valores[cel.coordinate] = "Página"
    cel = ws.cell(2, ULT)
    cel.number_format = "@"
    cel.alignment = _centro(False)
    cel = ws.cell(3, ULT, f"={capa_ref}!L3")
    cel.font = _fonte()
    cel.alignment = Alignment(horizontal="left", vertical="center")
    valores[cel.coordinate] = "CA"
    cel = ws.cell(4, ULT, f"={capa_ref}!L4")
    cel.alignment = _centro(False)
    valores[cel.coordinate] = capa.ca
    for r, campo in zip((5, 6, 7, 8), ("A5", "A6", "A7", "A8")):
        cel = ws.cell(r, 1, f"={capa_ref}!{campo}")
        cel.alignment = Alignment(vertical="center")
        valores[cel.coordinate] = {5: capa.obra, 6: capa.sistema, 7: capa.etapa, 8: capa.documento}[r]
    for r, h in [(1, 15), (2, 15), (3, 15), (4, 15), (5, 18), (6, 18), (7, 18), (8, 18), (9, 6)]:
        ws.row_dimensions[r].height = h
    if logo_cadari:
        img = XLImage(io.BytesIO(logo_cadari))
        alt = 57
        larg = round(alt * img.width / img.height)
        img.width, img.height = larg, alt
        col_c_px = config.LARGURAS_FIXAS["C"] * 7 + 5
        img.anchor = OneCellAnchor(
            _from=AnchorMarker(col=2, row=0, colOff=int(max(0, (col_c_px - larg) / 2)) * EMU, rowOff=11 * EMU),
            ext=XDRPositiveSize2D(cx=larg * EMU, cy=alt * EMU))
        ws.add_image(img)

    # ---------------- cabeçalho da tabela (linhas 10 a 13) ----------------
    for r in range(10, 14):
        for c in range(1, ULT + 1):
            cel = ws.cell(r, c)
            cel.fill = _preench(config.COR_CABECALHO)
            cel.font = _fonte(negrito=True)
            cel.alignment = _centro(True)
            cel.border = _borda(c, ULT)
    for ref, txt in (("A10", "Item"), ("B10", "Descrição"), ("D10", "Unid"), ("E10", "Total"),
                     ("D13", "Número de repetições")):
        ws[ref] = txt
    ws["D13"].alignment = _centro(False)
    ws.merge_cells("A10:A13")
    ws.merge_cells("B10:C13")
    ws.merge_cells("D10:D12")
    ws.merge_cells("E10:E12")
    ws.merge_cells("D13:E13")
    i = 0
    while i < n:
        j = i
        while j + 1 < n and colunas[j + 1].nivel == colunas[i].nivel:
            j += 1
        cfg = analise.niveis[colunas[i].nivel]
        ws.cell(10, PRIMEIRA + i, cfg.cabecalho)
        if j > i:
            ws.merge_cells(start_row=10, start_column=PRIMEIRA + i, end_row=10, end_column=PRIMEIRA + j)
        for k in range(i, j + 1):
            ws.cell(11, PRIMEIRA + k, colunas[k].titulo)
            ws.merge_cells(start_row=11, start_column=PRIMEIRA + k, end_row=12, end_column=PRIMEIRA + k)
            ws.cell(13, PRIMEIRA + k, cfg.repeticoes)
        i = j + 1
    for r, h in [(10, 27.0), (11, 27.0), (12, 43.0), (13, 20.1)]:
        ws.row_dimensions[r].height = h

    # ---------------- corpo ----------------
    def linha_titulo(r, codigo, texto, cor=config.COR_SECAO, h=20.1):
        for c in range(1, ULT + 1):
            cel = ws.cell(r, c)
            cel.fill = _preench(cor)
            cel.font = _fonte(negrito=True)
            cel.border = _borda(c, ULT)
            cel.alignment = _centro(True) if c == 1 else Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.cell(r, 1, codigo)
        ws.cell(r, 2, texto)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=ULT)
        ws.row_dimensions[r].height = h

    def linha_item(r, codigo, descricao, unidade, decimal):
        for c in range(1, ULT + 1):
            cel = ws.cell(r, c)
            cel.border = _borda(c, ULT)
            if c == 1:
                cel.font = _fonte(negrito=True)
                cel.alignment = _centro(True)
            elif c in (2, 3):
                cel.font = _fonte()
                cel.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif c in (4, 5):
                cel.font = _fonte(negrito=True)
                cel.alignment = _centro(True)
            else:
                cel.font = _fonte()
                cel.alignment = _centro(True)
                cel.number_format = "0.0" if decimal else "0"
        ws.cell(r, 1, codigo)
        ws.cell(r, 2, descricao)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.cell(r, 4, unidade)
        ws.cell(r, 5).number_format = "0" if decimal else "General"
        ws.row_dimensions[r].height = _altura_descricao(descricao)

    r = 14
    for c in range(1, ULT + 1):
        cel = ws.cell(r, c)
        cel.border = _borda(c, ULT, bottom=SEM)
        cel.font = _fonte(negrito=True)
        cel.alignment = _centro(True)
    ws.cell(r, 2, capa.sistema)
    ws.merge_cells("B14:C14")
    ws.row_dimensions[r].height = 19.5

    letras_secao = iter("ABCDEFGHIJKLMNOPQRSTUVWXYZ")
    for secao, subs in estrutura.items():
        presentes = []
        for sub, itens in subs.items():
            com_qtd = [it for it in itens if any(abs(q) > 1e-9 for q in it.qtd[aba].values())]
            if com_qtd:
                presentes.append((sub, com_qtd))
        if not presentes:
            continue
        L_sec = next(letras_secao)
        r += 1
        linha_titulo(r, L_sec, secao)
        n_sub = 0
        for sub, itens in presentes:
            if sub is not None:
                n_sub += 1
                r += 1
                linha_titulo(r, f"{L_sec}.{n_sub}", sub)
            for n_item, it in enumerate(itens, start=1):
                r += 1
                decimal = it.unidade.strip().lower() in config.UNIDADES_DECIMAIS
                codigo = f"{L_sec}.{n_sub}.{n_item}" if sub is not None else f"{L_sec}.{n_item}"
                linha_item(r, codigo, it.descricao, it.unidade, decimal)
                total = 0.0
                for lanc in lancar_linha(it, aba, colunas):
                    q = round(lanc.qtd, 2) if decimal else (int(round(lanc.qtd)) if abs(lanc.qtd - round(lanc.qtd)) < 1e-9
                                                            else round(lanc.qtd, 2))
                    total += q
                    cel = ws.cell(r, PRIMEIRA + lanc.inicio, q)
                    if lanc.fim > lanc.inicio:
                        ws.merge_cells(start_row=r, start_column=PRIMEIRA + lanc.inicio,
                                       end_row=r, end_column=PRIMEIRA + lanc.fim)
                        g.mesclagens += 1
                    if lanc.comentario:
                        cmt = Comment(lanc.comentario, "Lista de material")
                        cmt.width, cmt.height = 320, 110
                        cel.comment = cmt
                        g.comentarios += 1
                cel = ws.cell(r, 5, f"=SUM({letra(PRIMEIRA)}{r}:{L}{r})")
                valores[cel.coordinate] = round(total, 6)
                g.itens += 1

    # seções fixas: elétrica e outros
    L_ele = next(letras_secao)
    r += 1
    linha_titulo(r, L_ele, "ELÉTRICA")
    r += 1
    linha_item(r, f"{L_ele}.1", config.ITEM_ELETRICA, "vb.", False)
    ws.cell(r, 5, 1)
    ws.row_dimensions[r].height = 34.5
    L_out = next(letras_secao)
    r += 1
    linha_titulo(r, L_out, "OUTROS")
    for k, txt in enumerate(config.ITENS_OUTROS, start=1):
        r += 1
        linha_item(r, f"{L_out}.{k}", txt, "vb.", False)
        ws.cell(r, 5, 1)
        ws.row_dimensions[r].height = 24.95

    # notas
    notas = list(config.NOTAS)
    notas.append(config.NOTA_DESENHOS.format(desenhos=capa.desenhos.strip()) if capa.desenhos.strip()
                 else config.NOTA_DESENHOS_SEM_REF)
    r += 1
    for c in range(1, ULT + 1):
        ws.cell(r, c).border = Border(left=MEDIO if c == 1 else SEM, right=MEDIO if c == ULT else SEM)
    ws.cell(r, 1, "NOTAS:").font = _fonte(negrito=True)
    ws.row_dimensions[r].height = 20.1
    for k, txt in enumerate(notas):
        r += 1
        ultima = k == len(notas) - 1
        for c in range(1, ULT + 1):
            ws.cell(r, c).border = Border(left=MEDIO if c == 1 else SEM, right=MEDIO if c == ULT else SEM,
                                          bottom=MEDIO if ultima else SEM)
        cel = ws.cell(r, 1, txt)
        cel.font = _fonte()
        cel.alignment = Alignment(horizontal="left", vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ULT)
        ws.row_dimensions[r].height = 20.1
    g.ultima_linha = r

    # ---------------- impressão e exibição ----------------
    ws.print_area = f"A1:{L}{r}"
    ws.print_title_rows = "1:13"
    ws.row_breaks = RowBreak()
    ws.row_breaks.append(Break(id=r))
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    m = ws.page_margins
    m.left, m.right, m.top, m.bottom, m.header, m.footer = 0.0389, 0.0, 0.354, 0.0, 0.354, 0.315
    ws.oddHeader.right.text = "&P/&N"
    ws.oddHeader.right.font = "-,Negrito"
    ws.sheet_properties.tabColor = config.COR_ABA
    ws.sheet_view.showGridLines = False
    ws.sheet_view.showZeros = False
    ws.sheet_view.view = "pageBreakPreview"
    ws.sheet_view.zoomScale = 40 if n > 20 else 55
    ws.sheet_view.zoomScaleNormal = 70
    ws.sheet_view.tabSelected = False
    ws.freeze_panes = "F14"
    return g
